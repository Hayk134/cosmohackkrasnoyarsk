"""create_presentation.py

Generates a professional 16:9 widescreen PowerPoint presentation (.pptx)
for Team WHAT: Kosmo·MRV Satellite Carbon Verification Platform.

Structure:
- Slide 1: Название команды (WHAT) и название кейса.
- Slide 2: Проблема и рыночный вызов + Скриншот спутникового контроля нарушений (media_1789850446492.png).
- Slide 3: Решение Kosmo·MRV + Скриншот интерактивного гео-портала и аналитики участка (media_1789850446492.png).
- Slide 4: Пользователи и аудитория платформы + Скриншот кабинета землепользователя (media_1789851489908.png).
- Slide 5: Инвестиционная ценность и FinTech + Скриншот финансовой модели ROI/NPV (media_1789850961037.png).
- Slide 6: Климатические риски и устойчивость + Скриншот стресс-теста полигона 2024-2035 (media_1789850974683.png).
- Slide 7: Климатический AI-Анализ и отчеты + Скриншот AI-ассистента и рекомендаций (media_1789850966622.png).
- Slide 8: Исследование кейсов ТЗ (Тверь vs Мордовия) + Защита от гринвошинга на реальных данных.
- Slide 9: Архитектура, STAC API и 348 автотестов (100% PASS).
- Slide 10: Финал, команда WHAT, ссылки на GitHub Pages и GitVerse.

Style:
- Background: Dark Green (#0E2616 Deep Forest Dark Green)
- Cards / Containers: White (#FFFFFF)
- Text on Cards: Green (#14552A / #0E3C1E)
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# Color Palette: Dark Green Background, White Cards, Green Text
# -----------------------------------------------------------------------------
COLOR_BG = RGBColor(14, 38, 22)             # #0E2616 Deep Forest Dark Green
COLOR_CARD = RGBColor(255, 255, 255)        # #FFFFFF Pure Clean White Card
COLOR_CARD_BORDER = RGBColor(215, 235, 220) # #D7EBDC Soft subtle border
COLOR_PRIMARY_GREEN = RGBColor(20, 85, 42)  # #14552A Rich Dark Forest Green
COLOR_SECONDARY_GREEN = RGBColor(38, 128, 68)# #268044 Emerald Forest
COLOR_TEXT_MAIN = RGBColor(14, 60, 30)      # #0E3C1E Deep contrast Green for text on white
COLOR_TEXT_MUTED = RGBColor(60, 95, 75)     # #3C5F4B Muted green-gray for secondary text
COLOR_ACCENT_RED = RGBColor(185, 45, 45)    # #B92D2D Forest Berry Red (burns/alerts)
COLOR_ACCENT_AMBER = RGBColor(195, 115, 20) # #C37314 Warm Amber
COLOR_HEADER_WHITE = RGBColor(250, 255, 252)# #FAFFFC Crisp White for headings on dark green
COLOR_HEADER_MINT = RGBColor(160, 230, 185) # #A0E6B9 Mint for subheadings on dark green
COLOR_CHIP_BG = RGBColor(225, 245, 232)     # #E1F5E8 Soft Light Mint Chip

# Image assets
BASE_DIR = Path(__file__).resolve().parent
IMG_MAP = BASE_DIR / "media_1789850446492.png"
IMG_FINANCE = BASE_DIR / "media_1789850961037.png"
IMG_AI = BASE_DIR / "media_1789850966622.png"
IMG_STRESS = BASE_DIR / "media_1789850974683.png"
IMG_LANDOWNER = BASE_DIR / "media_1789851489908.png"


def apply_bg(slide):
    """Sets a solid deep forest dark green background for the slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG


def add_header(slide, title_text: str, category_text: str = "КОСМОХАКАТОН 2026 · КОМАНДА WHAT · KOSMO·MRV"):
    """Adds a standardized clean header section on dark green background."""
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_HEADER_MINT

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.7))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(23)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_HEADER_WHITE


def add_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=COLOR_CARD_BORDER):
    """Adds a rounded rectangle container card (White by default)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def add_slide_with_screenshot(slide, header_title: str, card_title: str, card_subtitle: str, 
                              bullets: list, img_path: Path, img_caption: str):
    """Template: Left white card with concise text, Right side with screenshot and white caption card."""
    apply_bg(slide)
    add_header(slide, header_title)

    # Left Card: Short explanation
    card_x = Inches(0.8)
    card_y = Inches(1.5)
    card_w = Inches(4.8)
    card_h = Inches(5.45)
    add_card(slide, card_x, card_y, card_w, card_h, COLOR_CARD, COLOR_CARD_BORDER)

    tb = slide.shapes.add_textbox(card_x + Inches(0.3), card_y + Inches(0.25), card_w - Inches(0.6), card_h - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = card_title
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY_GREEN

    if card_subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = card_subtitle.upper()
        p_sub.font.size = Pt(10.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_SECONDARY_GREEN
        p_sub.space_before = Pt(4)

    for item in bullets:
        pb = tf.add_paragraph()
        if isinstance(item, tuple):
            title, desc = item
            pb.text = f"• {title}: "
            pb.font.bold = True
            pb.font.size = Pt(12)
            pb.font.color.rgb = COLOR_TEXT_MAIN
            # Add normal text
            run = pb.add_run()
            run.text = desc
            run.font.bold = False
            run.font.size = Pt(12)
            run.font.color.rgb = COLOR_TEXT_MUTED
        else:
            pb.text = f"• {item}"
            pb.font.size = Pt(12)
            pb.font.color.rgb = COLOR_TEXT_MAIN
        pb.space_before = Pt(8)

    # Right Image Area
    img_x = Inches(5.85)
    img_y = Inches(1.5)
    img_w = Inches(6.68)
    img_h = Inches(4.75)

    if img_path.exists():
        # Outer card frame for picture
        add_card(slide, img_x, img_y, img_w, img_h, RGBColor(20, 25, 22), COLOR_CARD_BORDER)
        # Add picture slightly inset
        slide.shapes.add_picture(str(img_path), img_x + Inches(0.08), img_y + Inches(0.08), 
                                 width=img_w - Inches(0.16), height=img_h - Inches(0.16))

    # Caption card beneath image
    cap_y = Inches(6.32)
    cap_h = Inches(0.63)
    add_card(slide, img_x, cap_y, img_w, cap_h, COLOR_CARD, COLOR_CARD_BORDER)
    tb_c = slide.shapes.add_textbox(img_x + Inches(0.2), cap_y + Inches(0.1), img_w - Inches(0.4), cap_h - Inches(0.2))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    pc = tf_c.paragraphs[0]
    pc.text = f"📌 {img_caption}"
    pc.font.size = Pt(11)
    pc.font.bold = True
    pc.font.color.rgb = COLOR_PRIMARY_GREEN


def build_presentation(output_path: str = "Kosmo_MRV_Presentation_WHAT.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (Команда WHAT и название кейса)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_bg(s1)

    # Main decorative white card
    add_card(s1, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1), COLOR_CARD, COLOR_CARD_BORDER)

    # Team Chip
    chip = add_card(s1, Inches(1.6), Inches(1.8), Inches(3.2), Inches(0.46), COLOR_CHIP_BG, None)
    tf_c = chip.text_frame
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
    p_c = tf_c.paragraphs[0]
    p_c.text = "КОМАНДА: WHAT"
    p_c.alignment = PP_ALIGN.CENTER
    p_c.font.size = Pt(13)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_PRIMARY_GREEN

    # Case Title
    tb = s1.shapes.add_textbox(Inches(1.6), Inches(2.55), Inches(10.1), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "КЕЙС: ВЕРИФИКАЦИЯ ЛЕСОКЛИМАТИЧЕСКИХ ПРОЕКТОВ"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_GREEN

    p2 = tf.add_paragraph()
    p2.text = "Kosmo·MRV: Спутниковая платформа верификации поглощения углерода и выпуска углеродных единиц"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_SECONDARY_GREEN
    p2.space_before = Pt(12)

    # Chips row
    chips = [
        ("Хакатон", "КосмоХакатон 2026 (SR Data)"),
        ("Стандарт", "ГОСТ Р ИСО 14064-2"),
        ("Качество", "348 / 348 PASS (100%)"),
    ]
    chip_w = Inches(3.1)
    for i, (k, v) in enumerate(chips):
        cx = Inches(1.6) + i * (chip_w + Inches(0.4))
        add_card(s1, cx, Inches(4.8), chip_w, Inches(1.0), COLOR_CHIP_BG, COLOR_SECONDARY_GREEN)
        t_box = s1.shapes.add_textbox(cx + Inches(0.15), Inches(4.88), chip_w - Inches(0.3), Inches(0.85))
        t_frame = t_box.text_frame
        t_frame.word_wrap = True
        pk = t_frame.paragraphs[0]
        pk.text = k.upper()
        pk.font.size = Pt(9.5)
        pk.font.color.rgb = COLOR_TEXT_MUTED
        pv = t_frame.add_paragraph()
        pv.text = v
        pv.font.size = Pt(12.5)
        pv.font.bold = True
        pv.font.color.rgb = COLOR_PRIMARY_GREEN

    # =========================================================================
    # SLIDE 2: PROBLEM (Проблема + Скриншот контроля пожаров и нарушений)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_with_screenshot(
        s2,
        header_title="Проблема: Кризис доверия и скрытые потери углерода",
        card_title="В чем сложность верификации?",
        card_subtitle="Ключевые риски инвесторов и верификаторов",
        bullets=[
            ("Угроза гринвошинга", "Пожары и вырубки часто скрыты за облаками. Без спутникового аудита углеродные единицы продолжают продаваться даже на сгоревших участках."),
            ("Игнорирование формы Земли", "Плоские калькуляторы (100х100 м = 1 га) дают ошибку площади до 40% на широтах РФ из-за сжатия меридианов."),
            ("Занижение шума в 3-5 раз", "Слепое сложение погрешностей пикселей без учета пространственной корреляции искажает 95% доверительный интервал."),
            ("Дороговизна аудита", "Полевые комиссии стоят миллионы рублей и занимают от 3 месяцев на участок."),
        ],
        img_path=IMG_MAP,
        img_caption="Модуль контроля нарушений: автоматическое обнаружение очагов пожара (MODIS) и 5 100 пикс. потерь покрова (Hansen)"
    )

    # =========================================================================
    # SLIDE 3: SOLUTION (Решение Kosmo·MRV + Скриншот спутникового дашборда)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_with_screenshot(
        s3,
        header_title="Решение: Kosmo·MRV — Объективный математический арбитр",
        card_title="Спутниковый аудит в 1 клик",
        card_subtitle="14-ступенчатый конвейер МГЭИК и ГОСТ",
        bullets=[
            ("Геодезия WGS 84", "Интеграл меридиональной кривизны на эллипсоиде Земли. Субпиксельное взвешивание с погрешностью < 0.0003 га."),
            ("Мультисенсорный синтез", "Интеграция 4 спутниковых систем: ESA CCI Biomass v7.0, MODIS MCD64A1, Hansen GFC и Sentinel-2 L2A."),
            ("Моран 2D + VIF", "Оценка автокорреляции шума по 8 соседям (Queen) и расчет дисперсионного инфляционного фактора."),
            ("Защитные шлюзы (Safety Gates)", "Автоматический запрет выпуска (Q = 0), если лес деградировал или шум H/R >= 1.0."),
        ],
        img_path=IMG_MAP,
        img_caption="Интерактивная карта Kosmo·MRV: спутниковый слой Sentinel-2, теплокарта биомассы, динамика полога и метрики поглощения"
    )

    # =========================================================================
    # SLIDE 4: USERS (Пользователи + Скриншот кабинета Землепользователя)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_with_screenshot(
        s4,
        header_title="Пользователи платформы: Личный кабинет Землепользователя",
        card_title="Для кого создан сервис?",
        card_subtitle="Специализированные профили и роли",
        bullets=[
            ("Землепользователь / Девелопер", "Управляет участками, видит кадастровый статус, динамику прироста биомассы и квоты к выпуску (+15–20% монетизации)."),
            ("Верификатор / Аудитор", "Мгновенная независимая экспертиза за 1 минуту вместо 3 месяцев. Формульный водопад и неизменяемый аудит-лог."),
            ("Инвестор и ESG-Банк", "Оценка доходности и рисков, стресс-тестирование климатических шоков перед покупкой углеродных кредитов."),
            ("Оператор Реестра РФ", "Экспорт верификационного пакета по ГОСТ Р ИСО 14064-2 и проверка балансов."),
        ],
        img_path=IMG_LANDOWNER,
        img_caption="Интерфейс «Мои участки и мониторинг угодий»: кадастровый контур, динамика полога (NDVI 0.83), структура леса и доход 7.3 млн ₽"
    )

    # =========================================================================
    # SLIDE 5: FINTECH (Инвестиционный анализ + Скриншот финмодели ROI/NPV)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_with_screenshot(
        s5,
        header_title="FinTech: Инвестиционный анализ и финансовая окупаемость (ROI)",
        card_title="Обоснование капитала",
        card_subtitle="Дисконтированные потоки (DCF) на 15 лет",
        bullets=[
            ("Интерактивная модель DCF", "Учет CAPEX посадки (7 000 ₽/га), OPEX охраны и MRV (850 ₽/га), ставки дисконтирования WACC (12%)."),
            ("Расчет ключевых показателей", "Чистая приведенная стоимость (NPV: 39.5 млн ₽), норма доходности (IRR: 62%), окупаемость (Break-even: 2-й год)."),
            ("Сравнение сценариев", "Сравнение выгод сохранения леса и выпуска квот против сценария сплошной промышленной вырубки."),
            ("Мгновенный экспорт", "Формирование инвестиционного меморандума и PDF-отчета для кредитных комитетов банков в 1 клик."),
        ],
        img_path=IMG_FINANCE,
        img_caption="Модуль «Инвестиционный анализ»: график кумулятивного NPV, окупаемость на 2-й год, распределение денежных потоков"
    )

    # =========================================================================
    # SLIDE 6: STRESS-TEST (Стресс-тестирование рисков 2024-2035)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_with_screenshot(
        s6,
        header_title="Климатическая надежность: Стресс-тест полигона 2024–2035",
        card_title="Моделирование шоков",
        card_subtitle="Стандарты IPCC и буферный резерв",
        bullets=[
            ("5 сценариев катастроф", "Засуха 2027 г., лесной пожар, короед-типограф, шквальный ветровал и комплексный пирогенный шок."),
            ("Интерактивное управление риском", "Ползунки интенсивности катастрофы (10%–100%) и размера буферного пула (10%–30% по ГОСТ Р 58973)."),
            ("Проверка буферного пула", "Тест показывает, покроет ли страховой буфер 20% потерю углерода (-15 369 т CO2e) без списания квот инвестора."),
            ("Индекс устойчивости TCFD", "Расчет скоринга климатической устойчивости проекта для регуляторов и страховых компаний."),
        ],
        img_path=IMG_STRESS,
        img_caption="Модуль «Стресс-тест полигона»: буферный пул 20% выдержал шок, оценка сохранности квот инвестора и перерасчет Стресс-NPV"
    )

    # =========================================================================
    # SLIDE 7: CLIMATE AI (Климатический AI-Анализ и рекомендации)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_with_screenshot(
        s7,
        header_title="Климатический AI-Анализ: Умные рекомендации и прогноз",
        card_title="Интеллектуальный ассистент",
        card_subtitle="Модели IPCC CMIP6 + спутниковые прогнозы",
        bullets=[
            ("Персонализированная сводка", "Адаптация аналитики под 3 роли: Инвестор, Эколог, Землепользователь."),
            ("Прогноз рисков до 2027 г.", "Оценка вероятности засухи по моделям Росгидромета и международной программе CMIP6."),
            ("Защита доходов от квот", "Расчет объема поглощения (5 955 т CO2) и подтверждение неприкосновенности годового дохода (7.68 млн ₽)."),
            ("Практические шаги безопасности", "Автоматические предписания (создание минерализованных противопожарных полос от 1.4 м)."),
        ],
        img_path=IMG_AI,
        img_caption="Модуль «Климатический AI-Анализ»: сводка для полигона Мордовии с ответами на риски засухи и сохранение углеродного дохода"
    )

    # =========================================================================
    # SLIDE 8: RESEARCH CASE STUDY (Тверь vs Мордовия - Защита от гринвошинга)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_bg(s8)
    add_header(s8, "Исследование на реальных данных ТЗ: Тверь vs Мордовия")

    rw = Inches(5.74)
    rh = Inches(5.45)

    # Left: TVER
    add_card(s8, Inches(0.8), Inches(1.5), rw, rh, COLOR_CARD, COLOR_PRIMARY_GREEN)
    tb_tv = s8.shapes.add_textbox(Inches(1.05), Inches(1.75), rw - Inches(0.5), rh - Inches(0.5))
    tf_tv = tb_tv.text_frame
    tf_tv.word_wrap = True

    pt1 = tf_tv.paragraphs[0]
    pt1.text = "RU_TVER_01 · КОНТРОЛЬНЫЙ ЭТАЛОН"
    pt1.font.size = Pt(13)
    pt1.font.bold = True
    pt1.font.color.rgb = COLOR_PRIMARY_GREEN

    pt2 = tf_tv.add_paragraph()
    pt2.text = "Тверская область (1 750.5 га)"
    pt2.font.size = Pt(20)
    pt2.font.bold = True
    pt2.font.color.rgb = COLOR_TEXT_MAIN
    pt2.space_before = Pt(4)

    tver_facts = [
        "Потери древесного покрова Hansen GFC: 0 пикселей.",
        "Гари MODIS MCD64A1: 0 термальных аномалий.",
        "Устойчивый рост средней биомассы: 110.8 -> 115.1 т/га.",
        "Валовый эффект проекта R > 0 (превышение базовой линии).",
        "Отношение неопределенности: H/R = 0.18 (норматив выдержан).",
        "ИТОГ ВЕРИФИКАЦИИ: Выпуск Q = 2 814 единиц ОДОБРЕН.",
    ]
    for f in tver_facts:
        pf = tf_tv.add_paragraph()
        pf.text = f"• {f}"
        pf.font.size = Pt(12.5)
        pf.font.color.rgb = COLOR_TEXT_MAIN
        pf.space_before = Pt(8)

    # Right: MORDOVIA
    add_card(s8, Inches(6.78), Inches(1.5), rw, rh, COLOR_CARD, COLOR_ACCENT_RED)
    tb_mo = s8.shapes.add_textbox(Inches(7.03), Inches(1.75), rw - Inches(0.5), rh - Inches(0.5))
    tf_mo = tb_mo.text_frame
    tf_mo.word_wrap = True

    pm1 = tf_mo.paragraphs[0]
    pm1.text = "RU_MORDOVIA_03 · АВАРИЙНЫЙ ПОЛИГОН"
    pm1.font.size = Pt(13)
    pm1.font.bold = True
    pm1.font.color.rgb = COLOR_ACCENT_RED

    pm2 = tf_mo.add_paragraph()
    pm2.text = "Республика Мордовия (1 829.6 га)"
    pm2.font.size = Pt(20)
    pm2.font.bold = True
    pm2.font.color.rgb = COLOR_TEXT_MAIN
    pm2.space_before = Pt(4)

    mord_facts = [
        "Пожар августа 2021 г.: MODIS фиксирует 60 пикселей горения.",
        "Потери покрова Hansen GFC: более 5 100 пикселей.",
        "Спектральный индекс dNBR подтверждает сильное выгорание.",
        "Падение запаса: проектный сток ниже фонового тренда (R <= 0).",
        "Сработал автоматический защитный шлюз Gate 1.",
        "ИТОГ ВЕРИФИКАЦИИ: БЛОКИРОВКА ВЫПУСКА (Q = 0 ЕДИНИЦ).",
    ]
    for f in mord_facts:
        pf = tf_mo.add_paragraph()
        pf.text = f"• {f}"
        pf.font.size = Pt(12.5)
        pf.font.color.rgb = COLOR_TEXT_MAIN
        pf.space_before = Pt(8)

    # =========================================================================
    # SLIDE 9: ARCHITECTURE & RELIABILITY (Архитектура и тесты)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_bg(s9)
    add_header(s9, "Production-Ready архитектура и надежность системы")

    arch_blocks = [
        (
            "348 Автотестов (100% PASS)",
            "Полное покрытие пирамиды тестирования",
            "176 сквозных E2E-тестов (Tiers 1-4) выполняются за рекордные 0.28 секунды. 144 теста модулей API и 28 тестов краевых условий. Эталон ТЗ сошелся на 100%.",
            COLOR_PRIMARY_GREEN,
        ),
        (
            "Интеграция STAC API",
            "Открытые спутниковые каталоги",
            "Клиент к AWS Earth Search (sentinel-2-l2a) и Microsoft Planetary Computer (modis-64A1-061) с бесшовным локальным кэшем scene_metadata.json.",
            COLOR_SECONDARY_GREEN,
        ),
        (
            "Распределенный деплой",
            "GitVerse + GitHub Pages + Docker",
            "Frontend развернут на GitHub Pages с динамическим API_BASE. Backend контейнеризирован в Docker для любого VPS/Cloud. Исходный код на GitVerse.",
            COLOR_ACCENT_AMBER,
        ),
    ]

    aw = Inches(3.75)
    ah = Inches(5.45)
    for i, (title, sub, desc, col) in enumerate(arch_blocks):
        ax = Inches(0.8) + i * (aw + Inches(0.24))
        add_card(s9, ax, Inches(1.5), aw, ah, COLOR_CARD, COLOR_CARD_BORDER)

        tb = s9.shapes.add_textbox(ax + Inches(0.25), Inches(1.75), aw - Inches(0.5), ah - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = sub.upper()
        p2.font.size = Pt(10.5)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(12.5)
        p3.font.color.rgb = COLOR_TEXT_MAIN
        p3.space_before = Pt(14)

    # =========================================================================
    # SLIDE 10: CONCLUSION & CONTACTS (Финал)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_bg(s10)

    add_card(s10, Inches(0.8), Inches(1.0), Inches(11.733), Inches(5.5), COLOR_CARD, COLOR_PRIMARY_GREEN)

    tb_f = s10.shapes.add_textbox(Inches(1.4), Inches(1.4), Inches(10.5), Inches(4.7))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True

    pf1 = tf_f.paragraphs[0]
    pf1.text = "КОМАНДА: WHAT · ПРОЕКТ: KOSMO·MRV"
    pf1.font.size = Pt(15)
    pf1.font.bold = True
    pf1.font.color.rgb = COLOR_SECONDARY_GREEN

    pf2 = tf_f.add_paragraph()
    pf2.text = "Прозрачные, математически строгие углеродные инвестиции"
    pf2.font.size = Pt(28)
    pf2.font.bold = True
    pf2.font.color.rgb = COLOR_PRIMARY_GREEN
    pf2.space_before = Pt(8)

    pf3 = tf_f.add_paragraph()
    pf3.text = (
        "Kosmo·MRV решает проблему кризиса доверия в климатических проектах, объединяя геодезию эллипсоида Земли, "
        "многосенсорную спутниковую аналитику (Sentinel-2, MODIS, Hansen, ESA CCI), формулы ГОСТ Р ИСО 14064-2 и FinTech инструменты."
    )
    pf3.font.size = Pt(15)
    pf3.font.color.rgb = COLOR_TEXT_MUTED
    pf3.space_before = Pt(14)

    pf4 = tf_f.add_paragraph()
    pf4.text = (
        "📍 Исходный код на GitVerse: gitverse.ru/hackrus.experts/kosmo-krasnoiarsk_what_143\n"
        "🌐 Веб-сервис на GitHub Pages: hayk134.github.io/cosmohackkrasnoyarsk\n"
        "📑 Документация API: /docs (FastAPI Swagger UI)\n"
        "🏆 Решение кейса КосмоХакатона 2026 готово к промышленному внедрению"
    )
    pf4.font.size = Pt(13.5)
    pf4.font.bold = True
    pf4.font.color.rgb = COLOR_PRIMARY_GREEN
    pf4.space_before = Pt(22)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")


if __name__ == "__main__":
    build_presentation("c:/Users/gg/Desktop/shrek/Kosmo_MRV_Presentation_WHAT.pptx")
